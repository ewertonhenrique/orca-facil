"""Baixa a SINAPI mais recente publicada pela Caixa e gera os arquivos JSON do site.

Uso:
    python scripts/atualizar_sinapi.py            # procura o mês mais recente publicado
    python scripts/atualizar_sinapi.py 2026-09    # força um mês
    python scripts/atualizar_sinapi.py --arquivo SINAPI.zip --mes 2026-09   # usa um zip já baixado

Saída em dados/sinapi/:
    meta.json   mês de referência, data da atualização e UFs disponíveis
    <UF>.json   composições e insumos daquela UF, com custo não desonerado e desonerado
"""
import datetime as dt
import http.cookiejar
import io
import json
import re
import sys
import time
import unicodedata
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

import openpyxl

URL = "https://www.caixa.gov.br/Downloads/sinapi-relatorios-mensais/SINAPI-{ano}-{mes:02d}-formato-xlsx.zip"
SAIDA = Path(__file__).resolve().parent.parent / "dados" / "sinapi"
UFS = ["AC", "AL", "AM", "AP", "BA", "CE", "DF", "ES", "GO", "MA", "MG", "MS", "MT", "PA", "PB",
       "PE", "PI", "PR", "RJ", "RN", "RO", "RR", "RS", "SC", "SE", "SP", "TO"]
# Abas do arquivo de referência da Caixa (formato adotado a partir de 2025):
# C = composições, I = insumos; SD = sem desoneração, CD = com desoneração
ABAS = {"CSD": ("composicoes", "nd"), "CCD": ("composicoes", "d"),
        "ISD": ("insumos", "nd"), "ICD": ("insumos", "d")}


def norm(txt):
    txt = unicodedata.normalize("NFKD", str(txt or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", txt).strip().upper()


def baixar(ano, mes):
    url = URL.format(ano=ano, mes=mes)
    # o site da Caixa redireciona para gravar um cookie antes de liberar o arquivo
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
    opener.addheaders = [("User-Agent", "Mozilla/5.0 (Windows NT 10.0; Win64; x64) obim-orca"),
                         ("Accept", "*/*")]
    dados, erro = None, None
    for tentativa in range(3):
        try:
            with opener.open(url, timeout=180) as r:
                dados = r.read()
            break
        except urllib.error.HTTPError as e:
            erro = e
            if e.code == 404:
                break
        except Exception as e:
            erro = e
        time.sleep(5 * (tentativa + 1))
    if dados is None:  # 404 enquanto a Caixa não publica o mês
        aviso(f"{ano}-{mes:02d}: indisponível em {url} ({erro})")
        return None
    if not dados.startswith(b"PK"):
        aviso(f"{ano}-{mes:02d}: {url} não devolveu um zip (início: {dados[:80]!r})")
        return None
    print(f"  {ano}-{mes:02d}: baixado ({len(dados) // 1024} KB)")
    return dados


def aviso(msg):
    # aparece como anotação no resumo da execução do GitHub Actions
    print(f"::warning::{msg}", flush=True)


def abrir_planilha(zip_bytes):
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as z:
        nomes = [n for n in z.namelist() if n.lower().endswith(".xlsx")]
        # o arquivo principal tem "Refer" no nome; senão pega o maior
        nomes.sort(key=lambda n: ("REFER" not in norm(n), -z.getinfo(n).file_size))
        if not nomes:
            raise RuntimeError("Nenhum .xlsx dentro do zip")
        print(f"  lendo {nomes[0]}")
        return openpyxl.load_workbook(io.BytesIO(z.read(nomes[0])), read_only=True, data_only=True)


def achar_aba(wb, sigla):
    for nome in wb.sheetnames:
        if norm(nome) == sigla or norm(nome).startswith(sigla + " "):
            return wb[nome]
    return None


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = str(v).strip().replace(".", "").replace(",", ".")
    try:
        return round(float(s), 2)
    except ValueError:
        return None


def ler_aba(ws):
    """Devolve {codigo: (descricao, unidade, {UF: custo})}."""
    linhas = ws.iter_rows(values_only=True)
    cab = None
    for i, linha in enumerate(linhas):
        n = [norm(c) for c in linha]
        if any(c.startswith("CODIGO") for c in n) and any(c.startswith("DESCRI") for c in n):
            cab = n
            break
        if i > 30:
            break
    if cab is None:
        raise RuntimeError(f"Cabeçalho não encontrado na aba {ws.title}")
    col_cod = next(i for i, c in enumerate(cab) if c.startswith("CODIGO"))
    col_desc = next(i for i, c in enumerate(cab) if c.startswith("DESCRI"))
    col_un = next((i for i, c in enumerate(cab) if c.startswith("UNID")), None)
    col_uf = {}
    for i, c in enumerate(cab):
        if c in UFS and c not in col_uf:  # primeira coluna com o nome da UF é o custo
            col_uf[c] = i
    if not col_uf:
        for j, linha in enumerate(ws.iter_rows(min_row=1, max_row=12, values_only=True), 1):
            aviso(f"{ws.title} linha {j}: " + " | ".join(str(c)[:25] for c in linha[:40] if c is not None))
        raise RuntimeError(f"Colunas de UF não encontradas na aba {ws.title}")
    itens = {}
    for linha in linhas:
        if linha is None or len(linha) <= col_desc:
            continue
        cod = linha[col_cod]
        if cod is None or str(cod).strip() == "":
            continue
        cod = str(cod).strip().split(".")[0] if isinstance(cod, float) else str(cod).strip()
        desc = str(linha[col_desc] or "").strip()
        un = str(linha[col_un] or "").strip() if col_un is not None else ""
        custos = {uf: num(linha[i]) for uf, i in col_uf.items() if i < len(linha)}
        itens[cod] = (desc, un, custos)
    return itens


def gerar(wb, referencia):
    por_uf = {uf: {"composicoes": {}, "insumos": {}} for uf in UFS}
    for sigla, (tipo, regime) in ABAS.items():
        ws = achar_aba(wb, sigla)
        if ws is None:
            raise RuntimeError(f"Aba {sigla} não encontrada. Abas: {wb.sheetnames}")
        itens = ler_aba(ws)
        print(f"  aba {sigla}: {len(itens)} itens")
        for cod, (desc, un, custos) in itens.items():
            for uf, custo in custos.items():
                reg = por_uf[uf][tipo].setdefault(cod, [cod, desc, un, None, None])
                reg[3 if regime == "nd" else 4] = custo
    SAIDA.mkdir(parents=True, exist_ok=True)
    ufs_ok = []
    for uf, dados in por_uf.items():
        if not dados["composicoes"]:
            continue
        doc = {
            "referencia": referencia,
            "uf": uf,
            "campos": ["codigo", "descricao", "unidade", "custo_nao_desonerado", "custo_desonerado"],
            "composicoes": sorted(dados["composicoes"].values(), key=lambda r: r[1]),
            "insumos": sorted(dados["insumos"].values(), key=lambda r: r[1]),
        }
        (SAIDA / f"{uf}.json").write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        ufs_ok.append(uf)
    meta = {"base": "SINAPI", "referencia": referencia,
            "atualizado_em": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
            "fonte": "Caixa Econômica Federal / IBGE", "ufs": ufs_ok}
    (SAIDA / "meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"  gerado: {len(ufs_ok)} UFs, referência {referencia}")


def referencia_atual():
    try:
        return json.loads((SAIDA / "meta.json").read_text(encoding="utf-8")).get("referencia")
    except (OSError, ValueError):
        return None


def main(args):
    if "--arquivo" in args:
        arq = Path(args[args.index("--arquivo") + 1])
        mes = args[args.index("--mes") + 1]
        gerar(abrir_planilha(arq.read_bytes()), mes)
        return
    hoje = dt.date.today()
    if args:
        ano, mes = map(int, args[0].split("-"))
        candidatos = [(ano, mes)]
    else:  # a Caixa publica com ~1 mês de atraso: tenta do mês passado para trás
        candidatos = []
        a, m = hoje.year, hoje.month
        for _ in range(4):
            m -= 1
            if m == 0:
                a, m = a - 1, 12
            candidatos.append((a, m))
    atual = referencia_atual()
    for ano, mes in candidatos:
        ref = f"{ano}-{mes:02d}"
        if ref == atual and not args:
            print(f"Base já está no mês mais recente ({ref}). Nada a fazer.")
            return
        dados = baixar(ano, mes)
        if dados:
            gerar(abrir_planilha(dados), ref)
            return
    sys.exit("Nenhum mês recente disponível na Caixa.")


if __name__ == "__main__":
    try:
        main(sys.argv[1:])
    except SystemExit as e:
        if e.code not in (None, 0):
            print(f"::error::{e.code}", flush=True)
        raise
    except Exception as e:
        print(f"::error::{type(e).__name__}: {e}", flush=True)
        raise
